from pathlib import Path
import hashlib
import json
import re
import shutil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'outputs' / 'smr-review'
OUT.mkdir(parents=True, exist_ok=True)
data = json.loads((ROOT / 'content' / 'smr-counterbrief.json').read_text())
sources = {s['id']: s for s in data['sources']}
assert len(sources) == len(data['sources'])

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.72)
sec.left_margin = sec.right_margin = Inches(0.85)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.widow_control = True
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].paragraph_format.space_after = Pt(15)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
doc.styles['Heading 2'].font.size = Pt(12)
doc.core_properties.title = data['title']
doc.core_properties.subject = 'Evidence review of Darlington SMR claims and the supplied research brief'
doc.core_properties.author = 'Prosper'

NS_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
notes = etree.Element(qn('w:footnotes'), nsmap={'w': NS_W, 'r': NS_R})
for note_id, note_type, separator in [(-1, 'separator', 'separator'), (0, 'continuationSeparator', 'continuationSeparator')]:
    note = etree.SubElement(notes, qn('w:footnote'), {qn('w:id'): str(note_id), qn('w:type'): note_type})
    etree.SubElement(etree.SubElement(etree.SubElement(note, qn('w:p')), qn('w:r')), qn('w:' + separator))
note_part = Part(PackURI('/word/footnotes.xml'), 'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml', b'', doc.part.package)
doc.part.relate_to(note_part, RT.FOOTNOTES)
note_count = 0

def xml_run(parent, text, size=18):
    run = etree.SubElement(parent, qn('w:r'))
    props = etree.SubElement(run, qn('w:rPr'))
    etree.SubElement(props, qn('w:sz'), {qn('w:val'): str(size)})
    etree.SubElement(props, qn('w:color'), {qn('w:val'): '000000'})
    t = etree.SubElement(run, qn('w:t'))
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    return run

def add_link(paragraph, label, url):
    rel = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '174A75'); props.append(color)
    run.append(props)
    t = OxmlElement('w:t'); t.text = label; run.append(t)
    h.append(run); paragraph._p.append(h)

def cite(paragraph, refs):
    global note_count
    if not refs:
        return
    for key in refs:
        assert key in sources, key
    note_count += 1
    run = paragraph.add_run()
    element = OxmlElement('w:footnoteReference')
    element.set(qn('w:id'), str(note_count))
    run._r.append(element)
    run.font.superscript = True
    note = etree.SubElement(notes, qn('w:footnote'), {qn('w:id'): str(note_count)})
    p = etree.SubElement(note, qn('w:p'))
    pp = etree.SubElement(p, qn('w:pPr'))
    etree.SubElement(pp, qn('w:spacing'), {qn('w:after'): '50', qn('w:line'): '220', qn('w:lineRule'): 'auto'})
    nr = etree.SubElement(p, qn('w:r'))
    etree.SubElement(nr, qn('w:footnoteRef'))
    xml_run(p, ' ')
    for i, key in enumerate(refs):
        s = sources[key]
        if i:
            xml_run(p, '; ')
        label = s['publisher'].split(',')[0] + ', ' + s['title']
        if s['url']:
            href = etree.SubElement(p, qn('w:hyperlink'))
            href.set(qn('r:id'), note_part.relate_to(s['url'], RT.HYPERLINK, is_external=True))
            xml_run(href, label)
        else:
            xml_run(p, label + ' (supplied Word brief)')
        if key == 'asuega':
            xml_run(p, ', manuscript pp. 2 and 6')
        elif key == 'ontario':
            xml_run(p, ', p. 47')
        elif key == 'opg-budget':
            xml_run(p, ', MD&A p. 12')
        elif key == 'iea2026':
            xml_run(p, ', pp. 46–50, 52, 58')
    xml_run(p, '.')

def paragraph(text, refs=None, style=None):
    p = doc.add_paragraph(style=style)
    lead = re.match(r'^(Confirmed status:|Counter:|Why it matters:|Limit:)(.*)$', text, re.S)
    if lead:
        p.add_run(lead.group(1)).bold = True
        p.add_run(lead.group(2))
    else:
        p.add_run(text)
    cite(p, refs or [])
    return p

def table(block):
    t = doc.add_table(rows=1, cols=len(block['headers']))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [Inches(w * 6.8 / sum(block['widths'])) for w in block['widths']]
    for col, width in zip(t.columns, widths):
        col.width = width
    for cell, value, width in zip(t.rows[0].cells, block['headers'], widths):
        cell.width = width
        cell.text = value
    for row in block['rows']:
        cells = t.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = width
            cell.text = value
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    e.set(qn('w:' + key), val)
                borders.append(e)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), '105'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            pr.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E9ECEF' if i == 0 else 'FFFFFF'); pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    run.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

doc.add_paragraph(data['title'], 'Title')
md = ['# ' + data['title'], '']
for section in data['sections']:
    doc.add_heading(section['heading'], 1)
    md.extend(['## ' + section['heading'], ''])
    for block in section['blocks']:
        if block['type'] == 'p':
            paragraph(block['text'], block.get('refs'))
            refs = ''.join('[^' + key + ']' for key in block.get('refs', []))
            md.extend([block['text'] + refs, ''])
        elif block['type'] == 'list':
            for item in block['items']:
                paragraph(item, style='List Bullet')
                md.append('- ' + item)
            md.append('')
        elif block['type'] == 'table':
            table(block)
            md.append('| ' + ' | '.join(block['headers']) + ' |')
            md.append('| ' + ' | '.join(['---'] * len(block['headers'])) + ' |')
            md.extend('| ' + ' | '.join(row) + ' |' for row in block['rows'])
            md.append('')

doc.add_heading('Sources', 1)
md.extend(['## Sources', ''])
for i, s in enumerate(data['sources'], 1):
    p = doc.add_paragraph()
    p.add_run(f'{i}. {s["publisher"]}. ')
    if s['url']:
        add_link(p, s['title'], s['url'])
    else:
        p.add_run(s['title'])
    p.add_run('. ' + s['date'] + '. ' + s['locator'])
    for run in p.runs:
        run.font.size = Pt(10)
    title = f'[{s["title"]}]({s["url"]})' if s['url'] else s['title']
    md.extend([f'[^' + s['id'] + f']: {s["publisher"]}. {title}. {s["date"]}. {s["locator"]}', ''])

note_part._blob = etree.tostring(notes, xml_declaration=True, encoding='UTF-8', standalone=True)
# Strip inherited decorative paragraph rules without changing table borders.
for style in doc.styles:
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
for border in list(doc.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
doc.save(OUT / 'PROSPER-SMR-COUNTERS-AND-PROOF.docx')
(OUT / 'PROSPER-SMR-COUNTERS-AND-PROOF.md').write_text('\n'.join(md))
reply_section = next(s for s in data['sections'] if s['heading'] == 'Reply suitable for sending')
reply_text = '\n\n'.join(b['text'] for b in reply_section['blocks'] if not b.get('refs'))
reply_links = '\n'.join(f'{sources[k]["publisher"]}: {sources[k]["url"]}' for k in ['cnsc','opg','ontario','asuega'])
(OUT / 'REPLY-TO-ZACKARY.txt').write_text(reply_text + '\n\nSources\n\n' + reply_links + '\n')

original = Path(data['originalPath'])
preserved = ROOT / 'source' / 'smr-brief' / original.name
preserved.parent.mkdir(parents=True, exist_ok=True)
sha = hashlib.sha256(original.read_bytes()).hexdigest()
if preserved.exists():
    assert hashlib.sha256(preserved.read_bytes()).hexdigest() == sha
else:
    shutil.copy2(original, preserved)
assert hashlib.sha256(preserved.read_bytes()).hexdigest() == sha
checks = {
    'sections': len(data['sections']), 'sources': len(sources), 'footnotes': note_count,
    'sourceReferencesResolve': True, 'sourceOriginalUnchanged': True,
    'sourceSha256': sha, 'preservedSource': str(preserved),
    'costReductionFrom6_1To4_1Percent': round((6.1-4.1)/6.1*100, 1),
    'costReductionFrom7_7To4_1Percent': round((7.7-4.1)/7.7*100, 1),
    'illustrative300MW90PercentTWh': 300*8760*.9/1e6,
    'visualQa': 'pending render and page inspection', 'publication': 'none'
}
(OUT / 'verification.json').write_text(json.dumps(checks, indent=2) + '\n')
print(json.dumps(checks, indent=2))
